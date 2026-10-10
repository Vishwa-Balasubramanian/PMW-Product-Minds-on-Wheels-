"""Step 1 - Ingest. Read the raw BMW files into SQLite and record every data problem found.

Owner: Arjun (data and evidence). Equipment parsing is best-effort and the obvious next thing to improve.
"""
from __future__ import annotations

import re
import logging
import shutil
import subprocess
from collections import defaultdict

import openpyxl

from . import db
from .config import (POWERTRAIN, RAW, SALES_FILE, STUDY_FILE, VEHICLES, country_to_market)
from .scope import survey_attribute_scope

SEVERITY_RE = re.compile(r"^I had (?:a|an)\s*(very minor|minor|serious|very serious)?\s*problem", re.I)
COMPETITORS = {
    "Mercedes": r"mercedes|\bbenz\b|maybach",
    "Audi": r"\baudi\b",
    "Tesla": r"\btesla\b",
    "Lexus": r"\blexus\b",
    "Porsche": r"\bporsche\b",
    "Genesis": r"\bgenesis\b",
    "Volvo": r"\bvolvo\b",
    "Lucid": r"\blucid\b",
    "Cadillac": r"\bcadillac\b",
    "Volkswagen": r"\bvw\b|volkswagen",
    "Range Rover": r"range rover|land rover",
    "Polestar": r"\bpolestar\b",
}
COMPETITOR_RES = {k: re.compile(v, re.I) for k, v in COMPETITORS.items()}


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def ingest_vehicles(con) -> None:
    for code, v in VEHICLES.items():
        con.execute("INSERT INTO vehicles (code, name, short, image) VALUES (?,?,?,?)",
                    (code, v["name"], v["short"], v["image"]))


def ingest_feedback(con, code: str) -> dict:
    path = RAW / VEHICLES[code]["feedback_file"]
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header, rows = rows[0], rows[1:]
    idx = {h: i for i, h in enumerate(header)}

    comments: dict[str, dict] = {}
    labels: dict[str, list] = defaultdict(list)
    for r in rows:
        fid = r[idx["ID"]]
        if not fid:
            continue
        text = (r[idx["Customer Feedback"]] or "").strip()
        if fid not in comments:
            comments[fid] = {
                "source": r[idx["Source"]],
                "country": r[idx["Country"]],
                "powertrain": r[idx["Engine Type"]],
                "text": text,
            }
        topic = r[idx["Vfc level2 Name"]]
        if topic in (None, "no_class_found"):
            topic = None
        labels[fid].append((r[idx["Feedback Type"]], topic, r[idx["Vfc level3 Name"]]))

    seen_texts: dict[str, str] = {}
    duplicates = 0
    for fid, c in comments.items():
        m = SEVERITY_RE.match(c["text"]) if c["source"] == "Source C" else None
        severity = None
        if m:
            severity = (m.group(1) or "moderate").lower()
        comps = [k for k, rx in COMPETITOR_RES.items() if rx.search(c["text"])]
        real_labels = [lb for lb in labels[fid] if lb[1]]
        if c["text"] in seen_texts:
            duplicates += 1
        seen_texts.setdefault(c["text"], fid)
        con.execute(
            "INSERT INTO feedback (id, vehicle, source, country, market, powertrain, text, severity, competitors, classified)"
            " VALUES (?,?,?,?,?,?,?,?,?,?)",
            (fid, code, c["source"], c["country"], country_to_market(c["country"]), c["powertrain"], c["text"],
             severity, ",".join(comps) or None, 1 if real_labels else 0),
        )
        for ftype, topic, sub in real_labels:
            con.execute(
                "INSERT INTO feedback_labels (feedback_id, feedback_type, topic, subtopic) VALUES (?,?,?,?)",
                (fid, ftype, topic, sub),
            )

    n_rows, n_comments = len(rows), len(comments)
    unclassified = sum(1 for fid in comments if not any(lb[1] for lb in labels[fid]))
    if n_rows != n_comments:
        db.issue(con, "feedback", "warning",
                 f"{n_rows:,} rows but {n_comments:,} comments: one comment is split across several rows, one per label.",
                 "Stored one row per comment plus a separate labels table; comment IDs do not establish unique customers.", code)
    if unclassified:
        db.issue(con, "feedback", "warning",
                 f"{unclassified:,} comments ({unclassified / n_comments:.0%}) have no topic label.",
                 "Kept, marked unclassified. The AI classification step (signals.classify_unlabelled) labels them.", code)
    if duplicates:
        db.issue(con, "feedback", "info", f"{duplicates} comments repeat another comment's exact text.",
                 "Kept both; counted once per comment ID.", code)
    markets = defaultdict(int)
    for c in comments.values():
        markets[country_to_market(c["country"])] += 1
    for mk in ("CN",):
        if markets.get(mk, 0) == 0:
            db.issue(con, "feedback", "critical", f"No customer comments from {mk}.",
                     "Requirements for this market rely on survey means and web research only; confidence is lowered.", code)
    pts = {POWERTRAIN.get(c["powertrain"], c["powertrain"]) for c in comments.values()}
    db.audit(con, actor_type="system", actor="Pipeline", entity_type="dataset", entity_id=code, action="ingested feedback",
             after=f"{n_comments} comments, {n_rows} rows, powertrains: {', '.join(sorted(p for p in pts if p))}")
    return {"comments": n_comments, "rows": n_rows, "unclassified": unclassified}


def ingest_studies(con) -> None:
    wb = openpyxl.load_workbook(RAW / STUDY_FILE, read_only=True)

    # CN/EU: label row followed by a "Mean" row. Column order is fixed by the sheet header.
    cols = {1: ("G60", "G68", "CN"), 2: ("G70", None, "CN"), 3: ("F70", None, "EU"),
            4: ("G60", None, "EU"), 5: ("G70", None, "EU")}
    attr = None
    for r in wb["CN_EU_2025"].iter_rows(values_only=True):
        if r[0] == "Mean" and attr:
            for ci, (veh, variant, mk) in cols.items():
                val = _num(r[ci] if ci < len(r) else None)
                if val is None:
                    continue
                ok, why = survey_attribute_scope(attr)
                con.execute(
                    "INSERT INTO survey_scores (vehicle, variant, market, study, attribute, mean, scale, in_scope, scope_reason)"
                    " VALUES (?,?,?,?,?,?,?,?,?)", (veh, variant, mk, "CN_EU_2025", attr, val, 10, int(ok), why))
        elif r[0] and r[0] not in ("Dataset country",):
            attr = str(r[0]).strip()
    db.issue(con, "customer study", "info",
             "CN/EU study gives means only (1-10), with no sample sizes or distributions.",
             "Shown as means. Confidence treats a 1-point gap as meaningful but notes the missing sample size.")
    db.issue(con, "customer study", "info",
             "China figures for the 5 Series are for the long-wheelbase G68, not the G60.",
             "Stored with variant = G68 and labelled in the app.", "G60")

    # US: question row, "Sample total" row, then seven answer rows.
    order = ["I Hate It", "A Failure", "Unsatisfactory", "Satisfactory", "Excellent", "Delightful", "I Love It"]
    us_cols = {2: "G60", 3: "G70"}
    rows = list(wb["US_2025"].iter_rows(values_only=True))
    empty = []
    i = 0
    while i < len(rows):
        r = rows[i]
        if r[0] and i + 1 < len(rows) and rows[i + 1][0] == "Sample total":
            q = str(r[0]).strip()
            ans = {rows[j][0]: rows[j] for j in range(i + 2, min(i + 9, len(rows))) if rows[j][0] in order}
            for ci, veh in us_cols.items():
                vals = {k: _num(ans[k][ci]) for k in order if k in ans}
                if not vals or all(v is None for v in vals.values()):
                    empty.append(f"{q} ({veh})")
                    continue
                total = sum(v for v in vals.values() if v) or 1
                mean = sum((order.index(k) + 1) * (v or 0) for k, v in vals.items()) / total
                b3 = sum(vals.get(k) or 0 for k in order[:3]) / total
                t2 = sum(vals.get(k) or 0 for k in order[5:]) / total
                ok, why = survey_attribute_scope(q)
                con.execute(
                    "INSERT INTO survey_scores (vehicle, market, study, attribute, mean, scale, bottom3, top2, in_scope, scope_reason)"
                    " VALUES (?,?,?,?,?,?,?,?,?,?)", (veh, "US", "US_2025", q, mean, 7, b3, t2, int(ok), why))
            i += 9
        else:
            i += 1
    if empty:
        db.issue(con, "customer study", "warning",
                 f"{len(empty)} US survey questions have no answers: {', '.join(sorted(set(empty)))}.",
                 "Skipped. Requirements that would rely on them get no survey support.")
    db.audit(con, actor_type="system", actor="Pipeline", entity_type="dataset", entity_id="customer_studies",
             action="ingested customer studies", after="CN_EU_2025 means, US_2025 distributions")
    wb.close()


def ingest_sales(con) -> None:
    wb = openpyxl.load_workbook(RAW / SALES_FILE, read_only=True)
    for code, v in VEHICLES.items():
        ws = wb[v["sales_sheet"]]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        for r in rows[1:]:
            mk = str(r[1]).upper()
            mk = "ROW" if mk == "ROW" else mk
            for ci, h in enumerate(header):
                if not str(h).startswith("volume_"):
                    continue
                year = int(str(h).split("_")[1])
                raw = r[ci]
                if isinstance(raw, str) and raw.startswith("="):
                    db.issue(con, "sales", "warning",
                             f"{v['short']} {mk} {year}: cell holds a broken formula ({raw}) instead of a number.",
                             "Treated as 0, matching the market's other years.", code)
                    val = 0
                else:
                    val = int(_num(raw) or 0)
                con.execute("INSERT INTO sales (vehicle, market, year, volume) VALUES (?,?,?,?)", (code, mk, year, val))
    db.audit(con, actor_type="system", actor="Pipeline", entity_type="dataset", entity_id="sales",
             action="ingested sales volumes", after="2024, 2025, 2030 forecast by market")
    wb.close()


PRICE_RE = re.compile(r"^(?P<name>\S.*?)\s{2,}(?P<code>[0-9][0-9A-Z]{2}|[A-Z0-9]{3})\s+(?P<price>[\d.]+,[–-]|–,–)")
PKG_RE = re.compile(r"^(?P<name>[A-ZÄÖÜ][A-ZÄÖÜ &\-]+PAKET[A-ZÄÖÜ ]*)\s{2,}(?P<code>[0-9A-Z]{3})\s+(?P<price>[\d.]+,[–-])")


def ingest_equipment(con, code: str) -> None:
    """Best-effort parse of the German price list. Next step: proper package contents and per-trim columns."""
    pdf = RAW / VEHICLES[code]["options_file"]
    if shutil.which("pdftotext"):
        text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True,
                              text=True, encoding="utf-8", check=True).stdout
    else:
        # Cross-platform fallback: Windows builds also consume all supplied PDFs.
        from pypdf import PdfReader
        logger = logging.getLogger('pypdf._text_extraction._layout_mode._fixed_width_page')
        rotated = []
        def collect_rotation_warning(record):
            if record.getMessage().startswith('Rotated text discovered'):
                rotated.append(record.getMessage())
                return False
            return True
        logger.addFilter(collect_rotation_warning)
        reader = PdfReader(pdf)
        try:
            text = '\n'.join(p.extract_text(extraction_mode='layout') or '' for p in reader.pages)
        finally:
            reader.close()
            logger.removeFilter(collect_rotation_warning)
        if rotated:
            db.issue(con, 'equipment', 'warning', 'PDF tables contain rotated headings that layout extraction may omit.',
                     'Equipment entries are provisional; verify per-trim availability against the source PDF.', code)
    package = None
    n = 0
    for line in text.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        m = PKG_RE.match(s.strip())
        if m:
            package = m.group("name").strip()
            price = float(m.group("price").replace(".", "").split(",")[0])
            con.execute("INSERT INTO equipment (vehicle, name, code, price_eur, availability, package, raw) VALUES (?,?,?,?,?,?,?)",
                        (code, package.title(), m.group("code"), price, "optional", None, s.strip()))
            n += 1
            continue
        if s.strip().startswith("– ") and package:
            con.execute("INSERT INTO equipment (vehicle, name, code, price_eur, availability, package, raw) VALUES (?,?,?,?,?,?,?)",
                        (code, s.strip()[2:].strip(), None, None, "package_item", package.title(), s.strip()))
            n += 1
            continue
        if not s.startswith(" ") and s.isupper() is False and "■" in s and "□" not in s:
            name = re.split(r"\s{2,}", s.strip())[0]
            con.execute("INSERT INTO equipment (vehicle, name, code, price_eur, availability, package, raw) VALUES (?,?,?,?,?,?,?)",
                        (code, name, None, None, "standard", None, s.strip()))
            n += 1
            package = None
            continue
        m = PRICE_RE.match(s.strip())
        if m:
            price_txt = m.group("price")
            price = None if price_txt.startswith("–") else float(price_txt.replace(".", "").split(",")[0])
            con.execute("INSERT INTO equipment (vehicle, name, code, price_eur, availability, package, raw) VALUES (?,?,?,?,?,?,?)",
                        (code, m.group("name").strip(), m.group("code"), price, "optional", None, s.strip()))
            n += 1
            package = None
    db.issue(con, "equipment", "warning",
             "Price lists cover Germany only (July 2026). No US or China equipment data was provided.",
             "Equipment gaps for other markets are marked as an assumption.", code)
    db.audit(con, actor_type="system", actor="Pipeline", entity_type="dataset", entity_id=f"{code} price list",
             action="parsed equipment list", after=f"{n} entries (best-effort parser)")


def run(con) -> dict:
    ingest_vehicles(con)
    stats = {code: ingest_feedback(con, code) for code in VEHICLES}
    ingest_studies(con)
    ingest_sales(con)
    for code in VEHICLES:
        ingest_equipment(con, code)
    con.commit()
    return stats
